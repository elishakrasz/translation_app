"""Add addenda quotations, parsha/festival detection, chapter section, and Key ideas / Category columns
to sources/quotations-status.xlsx. Re-runnable: recomputes the auto columns, never overwrites the
translator's Key ideas / Category / Parsha (manual) entries."""
import openpyxl, re, glob
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
P = 'sources/quotations-status.xlsx'
PARSHIYOT = [("Bereshit","בראשית"),("Noah","נח"),("Lekh Lekha","לך לך"),("Vayera","וירא"),("Hayei Sarah","חיי שרה"),("Toldot","תולדות"),
 ("Vayetze","ויצא"),("Vayishlah","וישלח"),("Vayeshev","וישב"),("Miketz","מקץ"),("Vayigash","ויגש"),("Vayehi","ויחי"),("Shemot","שמות"),
 ("Va'era","וארא"),("Bo","בא"),("Beshalah","בשלח"),("Yitro","יתרו"),("Mishpatim","משפטים"),("Terumah","תרומה"),("Tezaveh","תצוה"),
 ("Ki Tisa","כי תשא"),("Vayakhel","ויקהל"),("Pekudei","פקודי"),("Vayikra","ויקרא"),("Zav","צו"),("Shemini","שמיני"),("Tazria","תזריע"),
 ("Mezora","מצורע"),("Ahrei Mot","אחרי מות"),("Kedoshim","קדושים"),("Emor","אמור"),("Behar","בהר"),("Behukotai","בחוקותי"),
 ("Bamidbar","במדבר"),("Naso","נשא"),("Beha'alotekha","בהעלותך"),("Shelah","שלח"),("Korah","קרח"),("Hukat","חקת"),("Balak","בלק"),
 ("Pinhas","פינחס"),("Matot","מטות"),("Masei","מסעי"),("Devarim","דברים"),("Va'ethanan","ואתחנן"),("Ekev","עקב"),("Re'eh","ראה"),
 ("Shoftim","שופטים"),("Ki Tetze","כי תצא"),("Ki Tavo","כי תבוא"),("Nizavim","נצבים"),("Vayelekh","וילך"),("Ha'azinu","האזינו"),("Vezot ha-Berakhah","וזאת הברכה")]
FESTIVALS = [("Rosh Hashanah","ראש השנה",r"Rosh ha-?Shanah|ראש השנה|ר\"ה"),("Yom Kippur","יום הכיפורים",r"Yom Kippur|Day of Atonement|יום הכפורים|יום הכיפורים|יוה\"כ"),
 ("Sukkot","סוכות",r"Sukkot|Succot|(?<![א-ת])[ובל]?סוכות(?![א-ת])"),("Shemini Atzeret / Simhat Torah","שמיני עצרת / שמחת תורה",r"Shemini Atzeret|Simhat Torah|שמיני עצרת|שמחת תורה"),
 ("Hanukkah","חנוכה",r"Hanukk?ah|Chanukah|חנוכה"),("Purim","פורים",r"Purim|(?<![א-ת])[ובל]?פורים(?![א-ת])"),("Pesach / Haggadah","פסח / הגדה",r"Pesach|Passover|Haggadah|(?<![א-ת])[ובל]?(?:פסח|הגדה)(?![א-ת])"),
 ("Shavuot","שבועות",r"Shavuot|(?<![א-ת])[ובל]?שבועות(?![א-ת])"),("Shabbat ha-Gadol","שבת הגדול",r"Shabbat ha-?Gadol|שבת הגדול"),("Shabbat Shuvah","שבת שובה",r"Shabbat Shuvah|שבת שובה"),
 ("Tisha be-Av","תשעה באב",r"Tisha be-?Av|Ninth of Av|תשעה באב|ט\"ב"),("Elul","אלול",r"\bElul\b|אלול"),("Shabbat","שבת",r"\bShabbat\b(?! ha-?Gadol| Shuvah)")]
VARIANTS = {"Lekh Lekha":["Lech Lecha","Lech-Lecha"],"Hayei Sarah":["Chayei Sara","Chayei Sarah","Hayye Sarah"],"Vayetze":["Vayetzei","Vayeitzei"],
 "Vayishlah":["Vayishlach"],"Vayehi":["Vayechi","Va-yehi"],"Beshalah":["Beshalach","Beshallah"],"Tezaveh":["Tetzaveh","Tetsaveh"],"Zav":["Tzav"],
 "Mezora":["Metzora","Metsora"],"Ahrei Mot":["Achrei Mot","Acharei Mot","Aharei Mot"],"Behukotai":["Bechukotai","Bechukosai"],
 "Beha'alotekha":["Beha'alotcha","Behaalotecha","Beha'alotkha"],"Shelah":["Sh'lach","Shlach","Shelach"],"Hukat":["Chukat","Hukkat"],
 "Pinhas":["Pinchas"],"Va'ethanan":["Vaetchanan","Va'etchanan"],"Ekev":["Eikev"],"Ki Tetze":["Ki Teitzei","Ki Tetzei","Ki Tetse"],
 "Nizavim":["Nitzavim","Nitsavim"],"Vayelekh":["Vayeilech","Vayelech"],"Ha'azinu":["Haazinu"],"Toldot":["Toldos","Toledot"],"Vayeshev":["Vayeishev"],
 "Miketz":["Mikeitz","Mikketz"],"Shemini":["Shmini"],"Tazria":["Tazri'a"],"Kedoshim":["K'doshim"],"Korah":["Korach"],"Re'eh":["Reeh"],"Shoftim":["Shofetim"],
 "Ki Tavo":["Ki Savo"],"Vezot ha-Berakhah":["V'Zot HaBerachah","Vezot Haberakhah"],"Bereshit":["Bereishit","Bereshis"],"Va'era":["Vaera"]}
AMBIG_HE = {"נח","בא","צו","ראה","אמור","בהר","שלח","עקב","דברים","שמות","תולדות","משפטים","קדושים","נשא","בלק","קרח","מסעי","בראשית","ויקרא","במדבר"}
AMBIG_EN = {"Bo","Noah","Zav","Emor","Behar","Ekev","Balak","Korah","Devarim","Shemot","Bereshit","Vayikra","Bamidbar","Naso"}
def detect(text):
    t = str(text or '')
    # book titles and Bible references that look like parsha names
    t = re.sub(r"תולדות (?:יעקב|אדם|אהרן|יצחק|הנפש)|Toldot (?:Ya'?akov|Yaakov|Adam|Aharon|Yitzhak)", " ", t)
    t = re.sub(r"(?<![א-ת])(?:בראשית|שמות|ויקרא|במדבר|דברים)\s+[א-ת]{1,3}[\s,:]", " ", t)
    t = re.sub(r"\b(?:Gen|Exod|Lev|Num|Deut)\.?\s*\d+", " ", t)
    found = []
    for en, he in PARSHIYOT:
        e = '(?:' + '|'.join(re.escape(x).replace("\\'", "'?").replace("\\ ", "[ -]") for x in [en] + VARIANTS.get(en, [])) + ')'
        if en in AMBIG_EN: pe = rf"(?:Parashat|Parshat|parashah|,)\s*{e}\b"
        else: pe = rf"\b{e}\b"
        ph = rf"(?:פרשת|פ'|,)\s*{he}(?![א-ת])" if he in AMBIG_HE else rf"(?<![א-ת]){he}(?![א-ת])"
        if re.search(pe, t) or re.search(ph, t): found.append(f"{he} ({en})")
    for en, he, pat in FESTIVALS:
        if re.search(pat, t): found.append(f"{he} ({en})")
    return '; '.join(dict.fromkeys(found))
def main():
    global SEC, FIRST
    # sections: (chapter, marker) -> heading in the English
    def sections():
        m = {}; first = {}
        for d in sorted(glob.glob('ch[01][0-9]-*_temp')):
            ch = int(d[2:4]); cur = ''
            for f in sorted(glob.glob(d + '/chunk[0-9][0-9][0-9][0-9].md')):
                n = f[-7:-3]; body = open(f, encoding='utf-8').read().split('## Notes')[0]
                first.setdefault((ch, n), None)
                for p in re.split(r'\n\s*\n', body):
                    h = re.match(r'^##\s+(.*)', p)
                    if h and not h.group(1).startswith('Notes'): cur = h.group(1).strip(); continue
                    if first[(ch, n)] is None: first[(ch, n)] = cur
                    for x in re.findall(r'\^(\d+)\^', p): m.setdefault((ch, int(x)), cur)
        return m, first
    SEC, FIRST = sections()
    wb = openpyxl.load_workbook(P); AQ = wb['All quotations']
    hdr = [c.value for c in AQ[1]]
    NEW = ['Chapter section (auto)', 'Parsha / festival (auto)', 'Parsha / festival (manual)', 'Key ideas', 'Category']
    for name in NEW:
        if name not in hdr:
            AQ.cell(1, len(hdr) + 1, name).font = Font(bold=True); hdr.append(name)
    col = {h: i + 1 for i, h in enumerate(hdr)}
    # --- add rows: addenda + ch11 nn35/39 (only once)
    existing = {(r[0], str(r[2]), str(r[3])) for r in AQ.iter_rows(min_row=2, values_only=True)}
    ADD_TITLES = {13: "Addendum A: Methodology", 14: "Addendum B: Historical Background", 15: "Addendum C: Contradictions",
                  16: "Addendum D: Enlightenment", 17: "Addendum E: Romanticism", 18: "Addendum F: Elitism"}
    HASIDIC = r"Siah Sarfei|Ramatayim|Kol Simhah|Torat Simhah|Or Simhah|Meir Einei|Me'ir Einei|Nifla'ot|Niflaot|Tiferet|Hiddushei|Emet ve-Emunah|Ohel|Kol Mevaser|Simhat Israel|Ma'amarei|Shem mi-Shmuel|Sefat Emet|Or ha-Gan|Hashavah|Eser|Siftei|Kedushat|Torat ha-Yehudi|Zekhuta|Imrei|Sefer Alexander|Likkutei|Keter|Amud ha-Emet|Divrei|Zera Kodesh|Noam|Toldot|Sippurei|Pardes|Mishmeret|Beit Ya'akov|Iggerot|Tzava'at|Ba'al Shem|Sihot|Netiv|Hesed le-Avraham|Bet"
    added = 0
    def notes_of(d):
        out = {}
        for f in sorted(glob.glob(d + '/chunk[0-9][0-9][0-9][0-9].md')):
            for m in re.finditer(r'(?m)^(\d+)\\\.\s*(.*)$', open(f, encoding='utf-8').read()): out[int(m.group(1))] = m.group(2)
        return out
    for ch in range(13, 19):
        d = glob.glob(f'ch{ch}-*_temp')[0]; nt = notes_of(d)
        for f in sorted(glob.glob(d + '/output_chunk[0-9][0-9][0-9][0-9].md')):
            n = f[-7:-3]; body = open(f, encoding='utf-8').read().split('## הערות')[0]
            for p in re.split(r'\n\s*\n', body):
                if not ((p.startswith('>') or '[לבדיקה' in p) and p.strip() != '>'): continue
                ms = re.findall(r'\^(\d+)\^', p); note = ','.join(ms) or '—'
                cited = '; '.join(nt.get(int(x), '')[:160] for x in ms) or '(no note — see surrounding text)'
                kind = 'hasidic' if re.search(HASIDIC, cited) else 'modern?'
                if (ch, note, n) in existing: continue
                AQ.append([ch, ADD_TITLES[ch], note, n, cited, kind, 'not-checked', 'Still translation',
                           'Addenda draft 2026-10-12, auto-extracted (block quote or [לבדיקה]); verify'] + [None] * (len(hdr) - 9))
                existing.add((ch, note, n)); added += 1
    for note, txt in (("35", "body block quote, no source named in n. 35 (Tiferet ha-Yehudi context) — identify"),
                      ("39", "body block quote; n. 39 cites Siah Sarfei Kodesh 4:39 §5; 3:134 §505 — verify")):
        if not any(k[0] == 11 and k[1] == note for k in existing):
            AQ.append([11, 'Reaction against Przysucha', note, '', txt, 'hasidic', 'not-checked', 'Still translation',
                       'found missing 2026-10-13 (block quote with marker, no row)'] + [None] * (len(hdr) - 9)); added += 1
    # --- auto columns for every row
    AL = {(r[0], str(r[4])): r[3] for r in wb["Author's Hebrew list"].iter_rows(min_row=2, values_only=True) if r[4]}
    for r in AQ.iter_rows(min_row=2):
        ch = r[0].value; notes = [int(x) for x in re.findall(r'\d+', re.sub(r'\(.*?\)', '', str(r[2].value)))]
        sec = next((SEC[(ch, x)] for x in notes if (ch, x) in SEC), None)
        if sec is None and r[3].value: sec = FIRST.get((ch, f'{int(str(r[3].value)):04d}' if str(r[3].value).isdigit() else ''), '')
        r[col['Chapter section (auto)'] - 1].value = sec or ''
        text = ' | '.join(str(x or '') for x in (r[4].value, r[8].value, AL.get((ch, str(r[2].value)), '')))
        r[col['Parsha / festival (auto)'] - 1].value = detect(text)
    # --- Lists sheet + validation
    if 'Lists' not in wb.sheetnames:
        L = wb.create_sheet('Lists')
        L.append(['Parsha / festival', 'Key ideas (fill in)', 'Categories (fill in)'])
        for c in L[1]: c.font = Font(bold=True)
        vals = [f"{he} ({en})" for en, he in PARSHIYOT] + [f"{he} ({en})" for en, he, _ in FESTIVALS]
        for i, v in enumerate(vals, 2): L.cell(i, 1, v)
        L.column_dimensions['A'].width = 34; L.column_dimensions['B'].width = 40; L.column_dimensions['C'].width = 40
        L.cell(1, 5, 'Fill columns B and C with your key ideas and categories; the dropdowns in "All quotations" read from here (up to row 200).')
    n = AQ.max_row
    for name, rng in (('Parsha / festival (manual)', "=Lists!$A$2:$A$200"), ('Key ideas', "=Lists!$B$2:$B$200"), ('Category', "=Lists!$C$2:$C$200")):
        L = openpyxl.utils.get_column_letter(col[name])
        dv = DataValidation(type='list', formula1=rng, allow_blank=True, showErrorMessage=False)
        AQ.add_data_validation(dv); dv.add(f'{L}2:{L}{n}')
    for name, w in zip(NEW, (30, 28, 28, 30, 22)): AQ.column_dimensions[openpyxl.utils.get_column_letter(col[name])].width = w
    AQ.auto_filter.ref = f'A1:{openpyxl.utils.get_column_letter(len(hdr))}{n}'
    # Summary chapter rows for the addenda
    S = wb['Summary']
    if not any(str(c.value).startswith('13') for c in S['A'] if c.value is not None):
        pass
    wb.save(P)
    auto = sum(1 for r in AQ.iter_rows(min_row=2, values_only=True) if r[col['Parsha / festival (auto)'] - 1])
    print('rows added', added, '| total rows', n - 1, '| rows with parsha/festival detected', auto)


if __name__ == "__main__":
    main()

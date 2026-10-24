"""Build sources/sources-overview.xlsx: the source texts the translation's originals come from,
the works cited (bibliography), the sources still needed, and what was taken from the author's own
"Hebrew Sources of the Main Quotations" list.

Usage: python3 tools/build_sources_overview.py      (re-run after any sources round)
Reads sources/quotations-status.xlsx and work/bibliography/reshimat-kitzurim.md; writes only the overview.
"""
import os
import re
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = (Path(os.environ["TRANSLATOR_PROJECT"]).expanduser().resolve() if os.environ.get("TRANSLATOR_PROJECT")
        else sys.exit("Set TRANSLATOR_PROJECT to the book folder, e.g. TRANSLATOR_PROJECT=~/the_translator"))
import ast  # noqa: E402
# HebrewBooks number → title, read from tools/find_source.py without running it
_src = (Path(__file__).resolve().parent / "find_source.py").read_text(encoding="utf-8")
TITLES = ast.literal_eval(re.search(r"^TITLES = (\{.*?\n\})", _src, re.S | re.M).group(1))

OUT = ROOT / "sources" / "sources-overview.xlsx"
CLASSICAL = re.compile(r"Genesis|Exodus|Leviticus|Lev\b|Numbers|Deuteronomy|Psalms|Proverbs|Isaiah|Jeremiah|Micah|Ecclesiastes|Kings|"
                       r"בראשית|שמות|ויקרא|במדבר|דברים|תהלים|תהילים|משלי|ישעיה|ירמיה|יחזקאל|קהלת|מיכה|מלכים|שמואל|איוב|הושע|ברכות|שבת|סוכה|מגילה|חגיגה|סנהדרין|אבות|רבה|משנה|בבלי|ירושלמי|זוהר|Zohar|"
                       r"Niddah|Sotah|Chullin|Pesachim|Job\b|Deut|Gen\.|Exod|Num\.|siddur|Mishneh Torah|Rambam|שו\"ע|רמב\"ם|"
                       r"Berakhot|Shabbat|Sukkah|Megillah|Chagigah|Menachot|Avodah Zarah|Sanhedrin|Pirkei Avot|Avot|Rabbah|Rashi|Mishnah")
# step 2 (2026-10-12): entries already original whose text was switched to / corrected from the author's version
STEP2 = {(4, "54"): "switched to the author's version", (4, "106"): "Yiddish + back-translated gloss replaced by the author's Hebrew",
         (4, "105"): "uncertain word settled from the author's version", (4, "100"): "reading corrected (לאחר זמן) after scan check",
         (5, "59"): "switched to the author's version", (5, "83"): "switched to the author's Yiddish reading",
         (8, "33"): "Yiddish replaced by the author's Hebrew version", (8, "55"): "back-translated sentence replaced by the original",
         (10, "16"): "OCR gap filled from the author's version", (11, "48"): "switched to the author's version",
         (11, "49"): "switched to the author's edition", (11, "86"): "switched to the author's edition"}


def origins(where):
    w = str(where or "")
    keys = [("HebrewBooks scan", f"hb {n}", f"{TITLES.get(n, '?')} (HebrewBooks {n})", f"sources/hb/{n}.txt") for n in re.findall(r"hb\s*(\d+)", w)]
    if "Tiferet" in w or "hayehudy" in w:
        keys.append(("PDF supplied by the translator", "tiferet", "תפארת היהודי (פיעטרקוב תרע\"ב)", "Tiferet/hayehudy01–03.pdf"))
    if "Keter" in w:
        keys.append(("PDF supplied by the translator", "keter", "כתר כהונה, א' צדרבוים (אודסה תרכ\"ו)", "keter.pdf"))
    if "Author's Hebrew list" in w:
        keys.append(("The author's own list", "author", "Hebrew Sources of the Main Quotations (Mickey10 pp. 33–51)", "books/chapters/ch19-hebrew-sources.md"))
    for m in re.finditer(r"Sefaria[,:]?\s*([A-Za-z' ]+?)(?:[,;\d]|$)", w):
        name = m.group(1).strip()
        kind = "Sefaria — classical" if CLASSICAL.search(name) else "Sefaria — Hasidic / later"
        keys.append((kind, f"sef:{name}", name, "sources/sefaria/json"))
    if not keys and CLASSICAL.search(w):
        keys.append(("Sefaria — classical", "classical-ref", "Bible / Talmud / Midrash / halakhah / liturgy (by reference)", "sources/sefaria/json"))
    if not keys:
        keys.append(("Other / unrecorded", "other:" + w[:60], w[:60] or "—", ""))
    return {k[1]: k for k in keys}.values()


def style(ws, widths):
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="44546A")
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def main():
    q = openpyxl.load_workbook(ROOT / "sources" / "quotations-status.xlsx", read_only=True)
    rows = list(q["All quotations"].iter_rows(values_only=True))
    hdr = rows[0]
    wb = openpyxl.Workbook()

    # 1. source texts used
    used = defaultdict(lambda: {"n": 0, "chapters": set(), "notes": []})
    meta = {}
    for r in rows[1:]:
        if r[7] != "Original in text":
            continue
        for kind, key, title, path in origins(r[8]):
            meta[key] = (kind, title, path)
            u = used[key]; u["n"] += 1; u["chapters"].add(r[0]); u["notes"].append(f"{r[0]}:{r[2]}")
    ws = wb.active; ws.title = "Source texts used"
    ws.append(["Source text", "Kind", "Local file", "Quotations supplied", "Chapters", "Chapter:note"])
    for key, u in sorted(used.items(), key=lambda kv: -kv[1]["n"]):
        kind, title, path = meta[key]
        ws.append([title, kind, path, u["n"], ", ".join(map(str, sorted(u["chapters"]))), "; ".join(u["notes"])])
    style(ws, {"A": 46, "B": 26, "C": 30, "D": 12, "E": 18, "F": 80})

    # 2. works cited (bibliography draft)
    ws = wb.create_sheet("Works cited")
    ws.append(["Short form (as in the notes)", "Full entry (draft)", "Script", "Complete?", "What is missing"])
    bib = ROOT / "work" / "bibliography" / "reshimat-kitzurim.md"
    for line in bib.read_text(encoding="utf-8").splitlines() if bib.exists() else []:
        m = re.match(r"^\s*[-*]?\s*\*\*(.+?)\*\*\s*[—–-]\s*(.*)$", line)
        if not m:
            continue
        short, full = m.group(1), m.group(2)
        miss = "; ".join(re.findall(r"\[לבדיקה:?\s*([^\]]*)\]", full))
        ws.append([short, full, "Hebrew" if re.search(r"[א-ת]", short) else "Latin", "no" if "[לבדיקה" in full else "yes", miss])
    style(ws, {"A": 34, "B": 80, "C": 9, "D": 10, "E": 50})

    # 3. sources still needed
    ws = wb.create_sheet("Sources still needed")
    ws.append(["Book (not in our scans)", "Quotations it would supply", "Chapter:note"])
    need = defaultdict(list)
    for r in list(q["Still needed"].iter_rows(values_only=True))[1:]:
        need[r[4] or "(no source named)"].append(f"{r[0]}:{r[1]}")
    for book, notes in sorted(need.items(), key=lambda kv: -len(kv[1])):
        ws.append([book, len(notes), "; ".join(notes)])
    style(ws, {"A": 50, "B": 14, "C": 80})

    # 4. the author's Hebrew list
    ws = wb.create_sheet("From author's Hebrew list")
    ws.append(["Chapter", "Note (author's list)", "Author's source reference", "Outcome", "Detail", "Author's Hebrew text"])
    fills = {"Inserted": "E2EFDA", "Already original": "DDEBF7", "Not inserted": "FCE4D6", "No quotation": "EDEDED"}
    for r in list(q["Author's Hebrew list"].iter_rows(values_only=True))[1:]:
        ch, n, text, ref, cat = r[0], str(r[1]), r[2], r[3], r[8] or ""
        ins = (r[10] or "") if len(r) > 10 else ""
        if cat.startswith(("NEW", "UPG")) or "applied" in ins:
            out, det = ("Inserted", cat) if ins.startswith("applied") else ("Not inserted", ins.replace("not applied — ", ""))
        elif cat.startswith("Matched by hand"):
            out, det = ("Inserted" if "applied" in ins else "No quotation to replace"), cat
        elif "matches" in cat:
            out, det = "Already original — matches", ""
        elif "differs" in cat:
            k = (ch, str(r[4]))
            out, det = ("Already original — changed to the author's text (step 2)", STEP2[k]) if k in STEP2 else ("Already original — reviewed, kept", "")
        else:
            out, det = "No quotation row for this note", cat
        ws.append([ch, n, ref, out, det, text])
        key = next((k for k in fills if out.startswith(k)), None)
        if key:
            ws.cell(ws.max_row, 4).fill = PatternFill("solid", fgColor=fills[key])
        ws.cell(ws.max_row, 6).alignment = Alignment(wrap_text=True, horizontal="right", readingOrder=2, vertical="top")
    style(ws, {"A": 8, "B": 9, "C": 34, "D": 44, "E": 44, "F": 90})

    wb.save(OUT)
    for s in wb.worksheets:
        print(f"{s.title}: {s.max_row - 1} rows")


if __name__ == "__main__":
    main()
